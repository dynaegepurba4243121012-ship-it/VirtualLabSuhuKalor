import streamlit as st
import pandas as pd
import numpy as np
import matplotlib.pyplot as plt
import io
from datetime import datetime

from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT


# ============================================================
# KONFIGURASI
# ============================================================

st.set_page_config(
    page_title="Virtual Lab Suhu dan Kalor",
    page_icon="🌡️",
    layout="wide",
    initial_sidebar_state="expanded"
)


# ============================================================
# CSS
# ============================================================

st.markdown("""
<style>

.stApp {
    background-color: #f5f7fb;
}

section[data-testid="stSidebar"] {
    background: linear-gradient(180deg, #071a35 0%, #0d2a52 100%);
}

section[data-testid="stSidebar"] * {
    color: white !important;
}

.main-title {
    font-size: 42px;
    font-weight: 800;
    color: #0b1f3a;
    margin-bottom: 5px;
}

.subtitle {
    font-size: 19px;
    color: #64748b;
    margin-bottom: 25px;
}

.hero {
    background: linear-gradient(135deg, #0b2a52, #1261a0);
    padding: 35px;
    border-radius: 22px;
    color: white;
    margin-bottom: 25px;
}

.hero h1 {
    color: white;
    font-size: 35px;
    margin-bottom: 10px;
}

.hero p {
    color: #e5f1ff;
    font-size: 18px;
}

.card {
    background: white;
    padding: 25px;
    border-radius: 18px;
    box-shadow: 0 4px 15px rgba(0,0,0,0.07);
    margin-bottom: 20px;
}

.info-box {
    background: #eaf4ff;
    border-left: 5px solid #1976d2;
    padding: 15px;
    border-radius: 10px;
    margin: 15px 0;
}

.analysis-box {
    background: #f0fdf4;
    border-left: 5px solid #16a34a;
    padding: 18px;
    border-radius: 10px;
    margin-top: 15px;
}

.question-box {
    background: white;
    border: 1px solid #dbe3ee;
    padding: 20px;
    border-radius: 14px;
    margin-bottom: 15px;
}

.footer {
    text-align: center;
    color: #64748b;
    padding: 30px;
}

</style>
""", unsafe_allow_html=True)


# ============================================================
# SESSION STATE
# ============================================================

if "halaman" not in st.session_state:
    st.session_state.halaman = "Beranda"

if "hasil_praktikum" not in st.session_state:
    st.session_state.hasil_praktikum = []


# ============================================================
# FUNGSI SIMPAN
# ============================================================

def simpan_hasil(data):
    st.session_state.hasil_praktikum.append(data)


# ============================================================
# FUNGSI WORD
# ============================================================

def atur_font_document(doc):
    """
    Mengatur seluruh font dokumen menjadi Times New Roman.
    """

    for paragraph in doc.paragraphs:
        for run in paragraph.runs:
            run.font.name = "Times New Roman"
            run.font.size = Pt(12)

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.font.name = "Times New Roman"
                        run.font.size = Pt(12)


def buat_word(data):

    doc = Document()

    # --------------------------------------------------------
    # Margin
    # --------------------------------------------------------

    section = doc.sections[0]

    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin = Cm(3)
    section.right_margin = Cm(2.5)

    # --------------------------------------------------------
    # JUDUL
    # --------------------------------------------------------

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    r = p.add_run("HASIL PRAKTIKUM")
    r.bold = True
    r.font.size = Pt(16)

    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER

    r2 = p2.add_run("VIRTUAL LAB SUHU DAN KALOR")
    r2.bold = True
    r2.font.size = Pt(14)

    doc.add_paragraph("")

    # --------------------------------------------------------
    # IDENTITAS
    # --------------------------------------------------------

    heading = doc.add_heading("A. Identitas Praktikan", level=1)
    heading.runs[0].font.name = "Times New Roman"

    tabel = doc.add_table(rows=3, cols=2)
    tabel.style = "Table Grid"
    tabel.alignment = WD_TABLE_ALIGNMENT.CENTER

    identitas = [
        ("Nama", data.get("Nama", "")),
        ("Kelas", data.get("Kelas", "")),
        ("Waktu", data.get("Waktu", ""))
    ]

    for i, (label, nilai) in enumerate(identitas):

        tabel.cell(i, 0).text = label
        tabel.cell(i, 1).text = str(nilai)

        tabel.cell(i, 0).vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        tabel.cell(i, 1).vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

    doc.add_paragraph("")

    # --------------------------------------------------------
    # JUDUL PRAKTIKUM
    # --------------------------------------------------------

    doc.add_heading("B. Judul Praktikum", level=1)
    doc.add_paragraph(str(data.get("Praktikum", "")))

    # --------------------------------------------------------
    # TUJUAN
    # --------------------------------------------------------

    doc.add_heading("C. Tujuan Pembelajaran", level=1)

    tujuan = data.get("Tujuan", "")

    for item in tujuan.split("\n"):
        if item.strip():
            doc.add_paragraph(item)

    # --------------------------------------------------------
    # LANGKAH
    # --------------------------------------------------------

    doc.add_heading("D. Langkah Praktikum", level=1)

    langkah = data.get("Langkah", "")

    for item in langkah.split("\n"):
        if item.strip():
            doc.add_paragraph(item)

    # --------------------------------------------------------
    # HASIL PENGAMATAN
    # --------------------------------------------------------

    doc.add_heading("E. Hasil Pengamatan", level=1)

    hasil_pengamatan = data.get("Data Utama", "")

    for item in hasil_pengamatan.split("\n"):
        if item.strip():
            doc.add_paragraph(item)

    # --------------------------------------------------------
    # HASIL PERHITUNGAN
    # --------------------------------------------------------

    doc.add_heading("F. Hasil Perhitungan", level=1)

    hasil = data.get("Hasil", "")

    for item in hasil.split("\n"):
        if item.strip():
            doc.add_paragraph(item)

    # --------------------------------------------------------
    # ANALISIS
    # --------------------------------------------------------

    doc.add_heading("G. Analisis", level=1)

    analisis = data.get("Analisis", "")

    doc.add_paragraph(analisis)

    # --------------------------------------------------------
    # PERTANYAAN DAN JAWABAN
    # --------------------------------------------------------

    doc.add_heading("H. Pertanyaan dan Jawaban", level=1)

    for i in range(1, 5):

        pertanyaan = data.get(f"Pertanyaan {i}", "")
        jawaban = data.get(f"Jawaban {i}", "")

        if pertanyaan:

            p = doc.add_paragraph()

            r = p.add_run(f"{i}. {pertanyaan}")
            r.bold = True

            p2 = doc.add_paragraph()

            r2 = p2.add_run("Jawaban: ")
            r2.bold = True

            p2.add_run(jawaban)

    # --------------------------------------------------------
    # KESIMPULAN
    # --------------------------------------------------------

    doc.add_heading("I. Kesimpulan", level=1)

    doc.add_paragraph(
        str(data.get("Kesimpulan", ""))
    )

    # --------------------------------------------------------
    # SIMPAN KE MEMORY
    # --------------------------------------------------------

    output = io.BytesIO()

    doc.save(output)

    output.seek(0)

    return output


# ============================================================
# SIDEBAR
# ============================================================

with st.sidebar:

    st.markdown(
        """
        <div style="
        text-align:center;
        padding:15px;
        ">
        <div style="font-size:50px;">🌡️</div>
        <h2>Virtual Lab</h2>
        <p>Suhu dan Kalor</p>
        </div>
        """,
        unsafe_allow_html=True
    )

    st.markdown("---")

    if st.button("🏠 Beranda", use_container_width=True):
        st.session_state.halaman = "Beranda"

    if st.button("🌡️ Praktikum Suhu", use_container_width=True):
        st.session_state.halaman = "Suhu"

    if st.button("🔥 Praktikum Kalor", use_container_width=True):
        st.session_state.halaman = "Kalor"

    if st.button("⚖️ Praktikum Asas Black", use_container_width=True):
        st.session_state.halaman = "Black"

    if st.button("💧 Praktikum Perubahan Wujud", use_container_width=True):
        st.session_state.halaman = "Wujud"

    if st.button("📋 Hasil Praktikum Saya", use_container_width=True):
        st.session_state.halaman = "Hasil"


# ============================================================
# BERANDA
# ============================================================

if st.session_state.halaman == "Beranda":

    st.markdown(
        """
        <div class="hero">

        <h1>Selamat Datang di VIRTUAL LAB SUHU DAN KALOR</h1>

        <p>
        Media Pembelajaran Praktikum Fisika Berbasis Web
        </p>

        <p>
        Materi: Suhu, Kalor, Asas Black, dan Perubahan Wujud
        </p>

        </div>
        """,
        unsafe_allow_html=True
    )

    st.markdown(
        """
        <div class="card">

        <h2>🔬 Tentang Virtual Lab</h2>

        <p>
        Virtual Lab Suhu dan Kalor merupakan media pembelajaran
        berbasis web yang memungkinkan peserta didik melakukan
        praktikum fisika secara interaktif.
        </p>

        <p>
        Setiap praktikum dilengkapi dengan tujuan pembelajaran,
        langkah praktikum, simulasi, pengamatan, analisis,
        pertanyaan, dan kesimpulan.
        </p>

        </div>
        """,
        unsafe_allow_html=True
    )

    st.subheader("🧪 Pilih Praktikum")

    col1, col2 = st.columns(2)

    with col1:

        st.markdown(
            """
            <div class="card">

            <h3>🌡️ Praktikum Suhu</h3>

            <p>
            Mempelajari pengukuran suhu dan hubungan
            antara beberapa skala suhu.
            </p>

            </div>
            """,
            unsafe_allow_html=True
        )

        if st.button(
            "Mulai Praktikum Suhu →",
            use_container_width=True
        ):
            st.session_state.halaman = "Suhu"
            st.rerun()

    with col2:

        st.markdown(
            """
            <div class="card">

            <h3>🔥 Praktikum Kalor</h3>

            <p>
            Mempelajari hubungan massa, kalor jenis,
            perubahan suhu, dan kalor.
            </p>

            </div>
            """,
            unsafe_allow_html=True
        )

        if st.button(
            "Mulai Praktikum Kalor →",
            use_container_width=True
        ):
            st.session_state.halaman = "Kalor"
            st.rerun()

    col3, col4 = st.columns(2)

    with col3:

        st.markdown(
            """
            <div class="card">

            <h3>⚖️ Praktikum Asas Black</h3>

            <p>
            Mempelajari pertukaran kalor hingga mencapai
            suhu keseimbangan.
            </p>

            </div>
            """,
            unsafe_allow_html=True
        )

        if st.button(
            "Mulai Praktikum Asas Black →",
            use_container_width=True
        ):
            st.session_state.halaman = "Black"
            st.rerun()

    with col4:

        st.markdown(
            """
            <div class="card">

            <h3>💧 Praktikum Perubahan Wujud</h3>

            <p>
            Mempelajari perubahan wujud zat akibat
            pemberian atau pelepasan kalor.
            </p>

            </div>
            """,
            unsafe_allow_html=True
        )

        if st.button(
            "Mulai Praktikum Perubahan Wujud →",
            use_container_width=True
        ):
            st.session_state.halaman = "Wujud"
            st.rerun()


# ============================================================
# PRAKTIKUM SUHU
# ============================================================

elif st.session_state.halaman == "Suhu":

    st.markdown(
        '<div class="main-title">🌡️ Praktikum Suhu</div>',
        unsafe_allow_html=True
    )

    st.markdown(
        '<div class="subtitle">Mengamati hubungan antar skala suhu</div>',
        unsafe_allow_html=True
    )

    nama = st.text_input("Nama Praktikan", key="nama_suhu")

    kelas = st.text_input("Kelas", key="kelas_suhu")

    st.markdown("### 🎯 Tujuan Pembelajaran")

    tujuan = """
1. Memahami konsep suhu sebagai besaran fisika.
2. Mengetahui hubungan antara skala Celsius, Fahrenheit, Kelvin, dan Reamur.
3. Mengamati perubahan nilai suhu pada beberapa skala.
"""

    st.info(tujuan)

    st.markdown("### 📋 Langkah Praktikum")

    langkah = """
1. Masukkan nama dan kelas.
2. Geser termometer pada nilai suhu yang diinginkan.
3. Amati perubahan nilai Celsius, Fahrenheit, Kelvin, dan Reamur.
4. Catat hasil pengamatan.
5. Jawab pertanyaan yang tersedia.
6. Buat kesimpulan berdasarkan hasil pengamatan.
"""

    st.write(langkah)

    st.markdown("---")

    st.subheader("🧪 Simulasi Pengukuran Suhu")

    suhu = st.slider(
        "Geser termometer untuk mengubah suhu",
        -50,
        150,
        25,
        key="slider_suhu"
    )

    fahrenheit = (9 / 5 * suhu) + 32
    kelvin = suhu + 273.15
    reamur = 4 / 5 * suhu

    col1, col2, col3, col4 = st.columns(4)

    col1.metric("Celsius", f"{suhu:.2f} °C")
    col2.metric("Fahrenheit", f"{fahrenheit:.2f} °F")
    col3.metric("Kelvin", f"{kelvin:.2f} K")
    col4.metric("Reamur", f"{reamur:.2f} °R")

    # --------------------------------------------------------
    # GAMBAR TERMOMETER
    # --------------------------------------------------------

    fig, ax = plt.subplots(figsize=(5, 5))

    ax.bar(
        [0],
        [suhu],
        width=0.5
    )

    ax.axhline(0)

    ax.set_ylim(-50, 150)

    ax.set_xlim(-1, 1)

    ax.set_ylabel("Suhu (°C)")

    ax.set_title("Visualisasi Perubahan Suhu")

    ax.set_xticks([])

    st.pyplot(fig)

    plt.close(fig)

    analisis_suhu = (
        f"Berdasarkan simulasi, suhu yang diamati adalah "
        f"{suhu:.2f} °C. Nilai tersebut setara dengan "
        f"{fahrenheit:.2f} °F, {kelvin:.2f} K, dan "
        f"{reamur:.2f} °R. Perubahan nilai pada setiap skala "
        f"terjadi karena setiap skala suhu mempunyai titik acuan "
        f"dan interval pembagian yang berbeda."
    )

    st.markdown("### 🔎 Analisis")

    st.markdown(
        f'<div class="analysis-box">{analisis_suhu}</div>',
        unsafe_allow_html=True
    )

    st.markdown("### ❓ Pertanyaan")

    q1 = st.text_input(
        "1. Apa yang dimaksud dengan suhu?",
        key="suhu_q1"
    )

    q2 = st.text_input(
        "2. Mengapa nilai suhu pada skala Celsius dan Fahrenheit berbeda?",
        key="suhu_q2"
    )

    q3 = st.text_input(
        "3. Berapa nilai Kelvin ketika suhu menunjukkan "
        f"{suhu:.2f} °C?",
        key="suhu_q3"
    )

    q4 = st.text_input(
        "4. Apa hubungan perubahan suhu dengan energi termal?",
        key="suhu_q4"
    )

    st.markdown("### 📝 Kesimpulan")

    kesimpulan = st.text_area(
        "Tuliskan kesimpulan praktikum:",
        key="suhu_kesimpulan"
    )

    if st.button(
        "📤 Kumpulkan Hasil Praktikum Suhu",
        use_container_width=True
    ):

        if not nama or not kelas:

            st.warning(
                "Silakan isi Nama dan Kelas terlebih dahulu."
            )

        else:

            data = {
                "Waktu": datetime.now().strftime("%d-%m-%Y %H:%M"),
                "Nama": nama,
                "Kelas": kelas,
                "Praktikum": "Praktikum Suhu",

                "Tujuan": tujuan,
                "Langkah": langkah,

                "Data Utama":
                    f"Suhu = {suhu:.2f} °C\n"
                    f"Fahrenheit = {fahrenheit:.2f} °F\n"
                    f"Kelvin = {kelvin:.2f} K\n"
                    f"Reamur = {reamur:.2f} °R",

                "Hasil":
                    f"Celsius = {suhu:.2f} °C\n"
                    f"Fahrenheit = {fahrenheit:.2f} °F\n"
                    f"Kelvin = {kelvin:.2f} K\n"
                    f"Reamur = {reamur:.2f} °R",

                "Analisis": analisis_suhu,

                "Pertanyaan 1":
                    "Apa yang dimaksud dengan suhu?",

                "Jawaban 1": q1,

                "Pertanyaan 2":
                    "Mengapa nilai suhu pada skala Celsius dan Fahrenheit berbeda?",

                "Jawaban 2": q2,

                "Pertanyaan 3":
                    f"Berapa nilai Kelvin ketika suhu menunjukkan {suhu:.2f} °C?",

                "Jawaban 3": q3,

                "Pertanyaan 4":
                    "Apa hubungan perubahan suhu dengan energi termal?",

                "Jawaban 4": q4,

                "Kesimpulan": kesimpulan
            }

            simpan_hasil(data)

            st.success(
                "Hasil praktikum berhasil dikumpulkan!"
            )


# ============================================================
# PRAKTIKUM KALOR
# ============================================================

elif st.session_state.halaman == "Kalor":

    st.markdown(
        '<div class="main-title">🔥 Praktikum Kalor</div>',
        unsafe_allow_html=True
    )

    st.markdown(
        '<div class="subtitle">Mengamati hubungan kalor dan perubahan suhu</div>',
        unsafe_allow_html=True
    )

    nama = st.text_input("Nama Praktikan", key="nama_kalor")

    kelas = st.text_input("Kelas", key="kelas_kalor")

    st.markdown("### 🎯 Tujuan Pembelajaran")

    tujuan = """
1. Memahami konsep kalor.
2. Mengetahui hubungan kalor dengan massa, kalor jenis, dan perubahan suhu.
3. Menghitung besar kalor yang diterima atau dilepaskan suatu benda.
"""

    st.info(tujuan)

    st.markdown("### 📋 Langkah Praktikum")

    langkah = """
1. Masukkan nama dan kelas.
2. Tentukan massa benda.
3. Tentukan kalor jenis benda.
4. Tentukan suhu awal dan suhu akhir.
5. Amati perubahan suhu pada grafik.
6. Hitung besar kalor.
7. Jawab pertanyaan dan buat kesimpulan.
"""

    st.write(langkah)

    st.markdown("---")

    st.subheader("🧪 Simulasi Kalor")

    col1, col2 = st.columns(2)

    with col1:

        massa = st.slider(
            "Massa benda (kg)",
            0.1,
            5.0,
            1.0,
            0.1,
            key="massa_kalor"
        )

        kalor_jenis = st.slider(
            "Kalor jenis (J/kg°C)",
            100,
            5000,
            4200,
            100,
            key="c_kalor"
        )

    with col2:

        suhu_awal = st.slider(
            "Suhu awal (°C)",
            0,
            100,
            25,
            key="awal_kalor"
        )

        suhu_akhir = st.slider(
            "Suhu akhir (°C)",
            0,
            150,
            75,
            key="akhir_kalor"
        )

    delta_t = suhu_akhir - suhu_awal

    Q = massa * kalor_jenis * delta_t

    st.markdown("### 📊 Hasil Simulasi")

    c1, c2, c3 = st.columns(3)

    c1.metric(
        "Perubahan Suhu",
        f"{delta_t:.2f} °C"
    )

    c2.metric(
        "Kalor",
        f"{Q:,.2f} J"
    )

    c3.metric(
        "Kalor",
        f"{Q / 1000:.2f} kJ"
    )

    # Grafik
    waktu = np.linspace(0, 10, 50)

    suhu_grafik = np.linspace(
        suhu_awal,
        suhu_akhir,
        50
    )

    fig, ax = plt.subplots(figsize=(8, 4))

    ax.plot(
        waktu,
        suhu_grafik,
        marker="o",
        markersize=3
    )

    ax.set_xlabel("Waktu relatif")
    ax.set_ylabel("Suhu (°C)")
    ax.set_title("Grafik Perubahan Suhu")

    ax.grid(True, alpha=0.3)

    st.pyplot(fig)

    plt.close(fig)

    if delta_t > 0:

        kondisi = (
            "Benda menerima kalor sehingga suhunya meningkat."
        )

    elif delta_t < 0:

        kondisi = (
            "Benda melepaskan kalor sehingga suhunya menurun."
        )

    else:

        kondisi = (
            "Tidak terjadi perubahan suhu sehingga kalor bersih "
            "yang dihitung bernilai nol."
        )

    analisis_kalor = (
        f"Massa benda sebesar {massa:.2f} kg dengan kalor jenis "
        f"{kalor_jenis:.2f} J/kg°C mengalami perubahan suhu "
        f"sebesar {delta_t:.2f} °C. Berdasarkan persamaan "
        f"Q = m × c × ΔT, besar kalor yang terlibat adalah "
        f"{Q:.2f} J atau {Q/1000:.2f} kJ. {kondisi}"
    )

    st.markdown("### 🔎 Analisis")

    st.markdown(
        f'<div class="analysis-box">{analisis_kalor}</div>',
        unsafe_allow_html=True
    )

    st.markdown("### ❓ Pertanyaan")

    q1 = st.text_input(
        "1. Apa yang dimaksud dengan kalor?",
        key="kalor_q1"
    )

    q2 = st.text_input(
        "2. Apa pengaruh massa benda terhadap kalor yang dibutuhkan?",
        key="kalor_q2"
    )

    q3 = st.text_input(
        "3. Apa yang terjadi jika perubahan suhu diperbesar?",
        key="kalor_q3"
    )

    q4 = st.text_input(
        "4. Tuliskan persamaan kalor dan jelaskan setiap besaran di dalamnya.",
        key="kalor_q4"
    )

    st.markdown("### 📝 Kesimpulan")

    kesimpulan = st.text_area(
        "Tuliskan kesimpulan praktikum:",
        key="kalor_kesimpulan"
    )

    if st.button(
        "📤 Kumpulkan Hasil Praktikum Kalor",
        use_container_width=True
    ):

        if not nama or not kelas:

            st.warning(
                "Silakan isi Nama dan Kelas terlebih dahulu."
            )

        else:

            data = {
                "Waktu": datetime.now().strftime("%d-%m-%Y %H:%M"),
                "Nama": nama,
                "Kelas": kelas,
                "Praktikum": "Praktikum Kalor",

                "Tujuan": tujuan,
                "Langkah": langkah,

                "Data Utama":
                    f"Massa = {massa:.2f} kg\n"
                    f"Kalor jenis = {kalor_jenis:.2f} J/kg°C\n"
                    f"Suhu awal = {suhu_awal:.2f} °C\n"
                    f"Suhu akhir = {suhu_akhir:.2f} °C\n"
                    f"Perubahan suhu = {delta_t:.2f} °C",

                "Hasil":
                    f"Q = m × c × ΔT\n"
                    f"Q = {massa:.2f} × {kalor_jenis:.2f} × {delta_t:.2f}\n"
                    f"Q = {Q:.2f} J\n"
                    f"Q = {Q/1000:.2f} kJ",

                "Analisis": analisis_kalor,

                "Pertanyaan 1":
                    "Apa yang dimaksud dengan kalor?",

                "Jawaban 1": q1,

                "Pertanyaan 2":
                    "Apa pengaruh massa benda terhadap kalor yang dibutuhkan?",

                "Jawaban 2": q2,

                "Pertanyaan 3":
                    "Apa yang terjadi jika perubahan suhu diperbesar?",

                "Jawaban 3": q3,

                "Pertanyaan 4":
                    "Tuliskan persamaan kalor dan jelaskan setiap besaran di dalamnya.",

                "Jawaban 4": q4,

                "Kesimpulan": kesimpulan
            }

            simpan_hasil(data)

            st.success(
                "Hasil praktikum berhasil dikumpulkan!"
            )


# ============================================================
# PRAKTIKUM ASAS BLACK
# ============================================================

elif st.session_state.halaman == "Black":

    st.markdown(
        '<div class="main-title">⚖️ Praktikum Asas Black</div>',
        unsafe_allow_html=True
    )

    st.markdown(
        '<div class="subtitle">Mengamati pertukaran kalor hingga tercapai keseimbangan</div>',
        unsafe_allow_html=True
    )

    nama = st.text_input("Nama Praktikan", key="nama_black")

    kelas = st.text_input("Kelas", key="kelas_black")

    st.markdown("### 🎯 Tujuan Pembelajaran")

    tujuan = """
1. Memahami konsep keseimbangan termal.
2. Menerapkan Asas Black pada pencampuran dua benda.
3. Menentukan suhu keseimbangan antara benda panas dan benda dingin.
"""

    st.info(tujuan)

    st.markdown("### 📋 Langkah Praktikum")

    langkah = """
1. Masukkan nama dan kelas.
2. Tentukan massa benda panas dan benda dingin.
3. Tentukan suhu awal masing-masing benda.
4. Atur kalor jenis kedua benda.
5. Amati suhu keseimbangan.
6. Bandingkan kalor yang dilepaskan dan diterima.
7. Jawab pertanyaan dan buat kesimpulan.
"""

    st.write(langkah)

    st.markdown("---")

    st.subheader("🧪 Simulasi Asas Black")

    col1, col2 = st.columns(2)

    with col1:

        st.markdown("### 🔥 Benda Panas")

        m1 = st.slider(
            "Massa benda panas (kg)",
            0.1,
            5.0,
            1.0,
            0.1,
            key="m1_black"
        )

        c1 = st.slider(
            "Kalor jenis benda panas",
            100,
            5000,
            4200,
            100,
            key="c1_black"
        )

        T1 = st.slider(
            "Suhu awal benda panas (°C)",
            30,
            150,
            80,
            key="T1_black"
        )

    with col2:

        st.markdown("### ❄️ Benda Dingin")

        m2 = st.slider(
            "Massa benda dingin (kg)",
            0.1,
            5.0,
            1.0,
            0.1,
            key="m2_black"
        )

        c2 = st.slider(
            "Kalor jenis benda dingin",
            100,
            5000,
            4200,
            100,
            key="c2_black"
        )

        T2 = st.slider(
            "Suhu awal benda dingin (°C)",
            0,
            50,
            20,
            key="T2_black"
        )

    Te = (
        (m1 * c1 * T1) +
        (m2 * c2 * T2)
    ) / (
        (m1 * c1) +
        (m2 * c2)
    )

    Q_lepas = m1 * c1 * (T1 - Te)

    Q_terima = m2 * c2 * (Te - T2)

    st.markdown("### 📊 Hasil Simulasi")

    col1, col2, col3 = st.columns(3)

    col1.metric(
        "Suhu Keseimbangan",
        f"{Te:.2f} °C"
    )

    col2.metric(
        "Kalor Dilepas",
        f"{Q_lepas:.2f} J"
    )

    col3.metric(
        "Kalor Diterima",
        f"{Q_terima:.2f} J"
    )

    # Grafik
    waktu = np.linspace(0, 10, 50)

    suhu_panas = np.linspace(T1, Te, 50)

    suhu_dingin = np.linspace(T2, Te, 50)

    fig, ax = plt.subplots(figsize=(8, 4))

    ax.plot(
        waktu,
        suhu_panas,
        label="Benda panas"
    )

    ax.plot(
        waktu,
        suhu_dingin,
        label="Benda dingin"
    )

    ax.axhline(
        Te,
        linestyle="--",
        label="Suhu keseimbangan"
    )

    ax.set_xlabel("Waktu relatif")
    ax.set_ylabel("Suhu (°C)")
    ax.set_title("Proses Menuju Keseimbangan Termal")

    ax.legend()
    ax.grid(True, alpha=0.3)

    st.pyplot(fig)

    plt.close(fig)

    analisis_black = (
        f"Benda panas memiliki suhu awal {T1:.2f} °C, sedangkan "
        f"benda dingin memiliki suhu awal {T2:.2f} °C. Setelah "
        f"terjadi pertukaran kalor, kedua benda mencapai suhu "
        f"keseimbangan sebesar {Te:.2f} °C. Benda panas melepaskan "
        f"kalor sebesar {Q_lepas:.2f} J, sedangkan benda dingin "
        f"menerima kalor sebesar {Q_terima:.2f} J. Hasil tersebut "
        f"menunjukkan prinsip Asas Black bahwa kalor yang dilepaskan "
        f"oleh benda bersuhu lebih tinggi sama dengan kalor yang "
        f"diterima benda bersuhu lebih rendah pada sistem ideal."
    )

    st.markdown("### 🔎 Analisis")

    st.markdown(
        f'<div class="analysis-box">{analisis_black}</div>',
        unsafe_allow_html=True
    )

    st.markdown("### ❓ Pertanyaan")

    q1 = st.text_input(
        "1. Apa yang dimaksud dengan keseimbangan termal?",
        key="black_q1"
    )

    q2 = st.text_input(
        "2. Mengapa benda panas mengalami penurunan suhu?",
        key="black_q2"
    )

    q3 = st.text_input(
        "3. Mengapa benda dingin mengalami kenaikan suhu?",
        key="black_q3"
    )

    q4 = st.text_input(
        "4. Jelaskan hubungan Asas Black dengan hukum kekekalan energi.",
        key="black_q4"
    )

    st.markdown("### 📝 Kesimpulan")

    kesimpulan = st.text_area(
        "Tuliskan kesimpulan praktikum:",
        key="black_kesimpulan"
    )

    if st.button(
        "📤 Kumpulkan Hasil Praktikum Asas Black",
        use_container_width=True
    ):

        if not nama or not kelas:

            st.warning(
                "Silakan isi Nama dan Kelas terlebih dahulu."
            )

        else:

            data = {
                "Waktu": datetime.now().strftime("%d-%m-%Y %H:%M"),
                "Nama": nama,
                "Kelas": kelas,
                "Praktikum": "Praktikum Asas Black",

                "Tujuan": tujuan,
                "Langkah": langkah,

                "Data Utama":
                    f"Massa benda panas = {m1:.2f} kg\n"
                    f"Suhu benda panas = {T1:.2f} °C\n"
                    f"Massa benda dingin = {m2:.2f} kg\n"
                    f"Suhu benda dingin = {T2:.2f} °C",

                "Hasil":
                    f"Suhu keseimbangan = {Te:.2f} °C\n"
                    f"Kalor dilepas = {Q_lepas:.2f} J\n"
                    f"Kalor diterima = {Q_terima:.2f} J",

                "Analisis": analisis_black,

                "Pertanyaan 1":
                    "Apa yang dimaksud dengan keseimbangan termal?",

                "Jawaban 1": q1,

                "Pertanyaan 2":
                    "Mengapa benda panas mengalami penurunan suhu?",

                "Jawaban 2": q2,

                "Pertanyaan 3":
                    "Mengapa benda dingin mengalami kenaikan suhu?",

                "Jawaban 3": q3,

                "Pertanyaan 4":
                    "Jelaskan hubungan Asas Black dengan hukum kekekalan energi.",

                "Jawaban 4": q4,

                "Kesimpulan": kesimpulan
            }

            simpan_hasil(data)

            st.success(
                "Hasil praktikum berhasil dikumpulkan!"
            )


# ============================================================
# PRAKTIKUM PERUBAHAN WUJUD
# ============================================================

elif st.session_state.halaman == "Wujud":

    st.markdown(
        '<div class="main-title">💧 Praktikum Perubahan Wujud</div>',
        unsafe_allow_html=True
    )

    st.markdown(
        '<div class="subtitle">Mengamati perubahan wujud zat akibat kalor</div>',
        unsafe_allow_html=True
    )

    nama = st.text_input("Nama Praktikan", key="nama_wujud")

    kelas = st.text_input("Kelas", key="kelas_wujud")

    st.markdown("### 🎯 Tujuan Pembelajaran")

    tujuan = """
1. Memahami berbagai jenis perubahan wujud zat.
2. Mengetahui hubungan kalor dengan perubahan wujud.
3. Mengamati perubahan susunan dan gerak partikel secara sederhana.
"""

    st.info(tujuan)

    st.markdown("### 📋 Langkah Praktikum")

    langkah = """
1. Masukkan nama dan kelas.
2. Pilih jenis perubahan wujud.
3. Atur energi/kalor menggunakan slider.
4. Amati perubahan visual partikel.
5. Amati perubahan yang terjadi pada zat.
6. Jawab pertanyaan.
7. Buat kesimpulan.
"""

    st.write(langkah)

    st.markdown("---")

    st.subheader("🧪 Simulasi Perubahan Wujud")

    proses = st.selectbox(
        "Pilih proses perubahan wujud",
        [
            "Mencair",
            "Membeku",
            "Menguap",
            "Mengembun",
            "Menyublim",
            "Mengkristal"
        ],
        key="proses_wujud"
    )

    energi = st.slider(
        "Atur energi/kalor",
        0,
        100,
        50,
        key="energi_wujud"
    )

    massa = st.slider(
        "Massa zat (kg)",
        0.1,
        5.0,
        1.0,
        0.1,
        key="massa_wujud"
    )

    kalor_laten = st.number_input(
        "Kalor laten (J/kg)",
        min_value=1.0,
        value=334000.0,
        step=1000.0,
        key="laten_wujud"
    )

    Q = massa * kalor_laten

    # --------------------------------------------------------
    # PARTIKEL
    # --------------------------------------------------------

    np.random.seed(10)

    jumlah = 80

    x = np.random.rand(jumlah) * 10
    y = np.random.rand(jumlah) * 6

    if proses in ["Mencair", "Menguap", "Menyublim"]:

        gerak = energi / 100 * 0.8

        x = x + np.random.randn(jumlah) * gerak
        y = y + np.random.randn(jumlah) * gerak

    elif proses in ["Membeku", "Mengembun", "Mengkristal"]:

        x = 5 + (x - 5) * (1 - energi / 150)
        y = 3 + (y - 3) * (1 - energi / 150)

    fig, ax = plt.subplots(figsize=(8, 5))

    ax.scatter(
        x,
        y,
        s=80
    )

    ax.set_xlim(0, 10)
    ax.set_ylim(0, 6)

    ax.set_title(
        f"Visualisasi Partikel - {proses}"
    )

    ax.set_xlabel("Posisi partikel")
    ax.set_ylabel("Posisi partikel")

    st.pyplot(fig)

    plt.close(fig)

    # --------------------------------------------------------
    # PENJELASAN
    # --------------------------------------------------------

    penjelasan = {

        "Mencair":
            "Mencair merupakan perubahan wujud dari padat menjadi cair "
            "karena zat menerima kalor.",

        "Membeku":
            "Membeku merupakan perubahan wujud dari cair menjadi padat "
            "karena zat melepaskan kalor.",

        "Menguap":
            "Menguap merupakan perubahan wujud dari cair menjadi gas "
            "karena zat menerima energi.",

        "Mengembun":
            "Mengembun merupakan perubahan wujud dari gas menjadi cair "
            "karena zat melepaskan kalor.",

        "Menyublim":
            "Menyublim merupakan perubahan wujud dari padat menjadi gas "
            "tanpa melalui fase cair.",

        "Mengkristal":
            "Mengkristal merupakan perubahan wujud dari gas menjadi padat."
    }

    st.markdown("### 🔎 Analisis")

    analisis_wujud = (
        f"Pada simulasi dipilih proses {proses} dengan energi "
        f"sebesar {energi}%. {penjelasan[proses]} "
        f"Jika digunakan massa {massa:.2f} kg dan kalor laten "
        f"{kalor_laten:,.0f} J/kg, maka energi kalor yang dihitung "
        f"dengan Q = mL adalah {Q:,.2f} J. "
        f"Perubahan wujud berkaitan dengan perubahan energi internal "
        f"dan susunan partikel zat."
    )

    st.markdown(
        f'<div class="analysis-box">{analisis_wujud}</div>',
        unsafe_allow_html=True
    )

    st.markdown("### ❓ Pertanyaan")

    q1 = st.text_input(
        "1. Apa yang dimaksud dengan perubahan wujud zat?",
        key="wujud_q1"
    )

    q2 = st.text_input(
        "2. Apa perbedaan mencair dan membeku?",
        key="wujud_q2"
    )

    q3 = st.text_input(
        "3. Mengapa perubahan wujud membutuhkan atau melepaskan kalor?",
        key="wujud_q3"
    )

    q4 = st.text_input(
        "4. Jelaskan hubungan gerak partikel dengan energi yang diberikan.",
        key="wujud_q4"
    )

    st.markdown("### 📝 Kesimpulan")

    kesimpulan = st.text_area(
        "Tuliskan kesimpulan praktikum:",
        key="wujud_kesimpulan"
    )

    if st.button(
        "📤 Kumpulkan Hasil Praktikum Perubahan Wujud",
        use_container_width=True
    ):

        if not nama or not kelas:

            st.warning(
                "Silakan isi Nama dan Kelas terlebih dahulu."
            )

        else:

            data = {
                "Waktu": datetime.now().strftime("%d-%m-%Y %H:%M"),
                "Nama": nama,
                "Kelas": kelas,
                "Praktikum": "Praktikum Perubahan Wujud",

                "Tujuan": tujuan,
                "Langkah": langkah,

                "Data Utama":
                    f"Proses = {proses}\n"
                    f"Energi = {energi}%\n"
                    f"Massa = {massa:.2f} kg\n"
                    f"Kalor laten = {kalor_laten:,.0f} J/kg",

                "Hasil":
                    f"Q = m × L\n"
                    f"Q = {massa:.2f} × {kalor_laten:,.0f}\n"
                    f"Q = {Q:,.2f} J",

                "Analisis": analisis_wujud,

                "Pertanyaan 1":
                    "Apa yang dimaksud dengan perubahan wujud zat?",

                "Jawaban 1": q1,

                "Pertanyaan 2":
                    "Apa perbedaan mencair dan membeku?",

                "Jawaban 2": q2,

                "Pertanyaan 3":
                    "Mengapa perubahan wujud membutuhkan atau melepaskan kalor?",

                "Jawaban 3": q3,

                "Pertanyaan 4":
                    "Jelaskan hubungan gerak partikel dengan energi yang diberikan.",

                "Jawaban 4": q4,

                "Kesimpulan": kesimpulan
            }

            simpan_hasil(data)

            st.success(
                "Hasil praktikum berhasil dikumpulkan!"
            )


# ============================================================
# HASIL PRAKTIKUM
# ============================================================

elif st.session_state.halaman == "Hasil":

    st.markdown(
        '<div class="main-title">📋 Hasil Praktikum Saya</div>',
        unsafe_allow_html=True
    )

    st.markdown(
        '<div class="subtitle">Hasil praktikum yang telah dikumpulkan</div>',
        unsafe_allow_html=True
    )

    if not st.session_state.hasil_praktikum:

        st.info(
            "Belum ada hasil praktikum yang dikumpulkan."
        )

    else:

        st.success(
            f"Terdapat {len(st.session_state.hasil_praktikum)} "
            "hasil praktikum pada sesi ini."
        )

        # ----------------------------------------------------
        # TAMPILKAN SATU-SATU
        # ----------------------------------------------------

        for nomor, data in enumerate(
            reversed(st.session_state.hasil_praktikum),
            start=1
        ):

            with st.expander(
                f"📄 {data['Praktikum']} - {data['Nama']}"
            ):

                st.write(
                    f"**Nama:** {data['Nama']}"
                )

                st.write(
                    f"**Kelas:** {data['Kelas']}"
                )

                st.write(
                    f"**Waktu:** {data['Waktu']}"
                )

                st.write(
                    f"**Praktikum:** {data['Praktikum']}"
                )

                st.markdown("### Hasil")

                st.text(
                    data["Hasil"]
                )

                st.markdown("### Analisis")

                st.write(
                    data["Analisis"]
                )

                st.markdown("### Kesimpulan")

                st.write(
                    data["Kesimpulan"]
                )

                # Word per hasil
                file_word = buat_word(data)

                nama_file = (
                    "Hasil_"
                    + data["Praktikum"].replace(" ", "_")
                    + "_"
                    + data["Nama"].replace(" ", "_")
                    + ".docx"
                )

                st.download_button(
                    "📄 Download Hasil Ini sebagai Word",
                    data=file_word,
                    file_name=nama_file,
                    mime=(
                        "application/vnd.openxmlformats-officedocument."
                        "wordprocessingml.document"
                    ),
                    key=f"download_word_{nomor}",
                    use_container_width=True
                )

        # ----------------------------------------------------
        # DOWNLOAD SEMUA
        # ----------------------------------------------------

        st.markdown("---")

        st.subheader("📄 Download Hasil Praktikum")

        st.write(
            "Setiap hasil praktikum dapat diunduh dalam bentuk "
            "dokumen Word yang sudah tersusun rapi."
        )


# ============================================================
# FOOTER
# ============================================================

st.markdown(
    """
    <div class="footer">

    <hr>

    <p>
    🌡️ <b>Virtual Lab Suhu dan Kalor</b>
    </p>

    <p>
    Media Pembelajaran Praktikum Fisika Berbasis Web
    </p>

    </div>
    """,
    unsafe_allow_html=True
)
